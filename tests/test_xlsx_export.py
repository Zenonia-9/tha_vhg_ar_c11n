from copy import deepcopy
from io import BytesIO
import json
from unittest.mock import patch

from openpyxl import load_workbook

from odoo.addons.account.tests.common import AccountTestInvoicingCommon
from odoo.addons.account_reports.models.account_report import AccountReport as NativeAccountReport
from odoo.tests import tagged


@tagged("post_install", "-at_install")
class TestAgedReceivableXlsxExport(AccountTestInvoicingCommon):
    def test_xlsx_matches_native_export(self):
        invoice = self.init_invoice(
            "out_invoice", invoice_date="2026-10-23", amounts=[100.0],
            taxes=self.env["account.tax"], post=True,
        )
        invoice.write({"anzer_id": "AR-XLSX-TEST", "vendor_ref": "Export reference"})
        report = self.env.ref("account_reports.aged_receivable_report")

        def workbook_cells(result):
            self.assertEqual(result["file_type"], "xlsx")
            self.assertTrue(result["file_name"].endswith(".xlsx"))
            workbook = load_workbook(BytesIO(result["file_content"]))
            try:
                return {
                    sheet.title: [
                        [(cell.value, cell.data_type, cell.number_format) for cell in row]
                        for row in sheet.iter_rows()
                    ]
                    for sheet in workbook.worksheets
                }
            finally:
                workbook.close()

        for unfold_all in (False, True):
            for show_currency in (False, True):
                options = report.get_options({
                    "date": {"date_from": False, "date_to": "2026-10-23", "mode": "single", "filter": "custom"},
                    "unfold_all": unfold_all, "show_currency": show_currency,
                })
                action = report.dispatch_report_action(options, "export_file", "export_to_xlsx", True)
                self.assertEqual(action["data"]["file_generator"], "export_to_xlsx")
                download_options = json.loads(action["data"]["options"])
                self.assertEqual(download_options["export_mode"], "file")
                with patch.object(type(report), "_build_static_line_columns", NativeAccountReport._build_static_line_columns):
                    expected = workbook_cells(report.export_to_xlsx(deepcopy(download_options)))
                actual = workbook_cells(report.export_to_xlsx(deepcopy(download_options)))
                self.assertEqual(actual, expected)
                cells = [cell[0] for rows in actual.values() for row in rows for cell in row]
                self.assertIn("Anzer ID", cells)
                self.assertIn("Invoice No.", cells)
                if unfold_all:
                    self.assertIn("AR-XLSX-TEST", cells)
                    self.assertIn("Export reference", cells)
                    self.assertIn(invoice.name, cells)
                    self.assertIn(100.0, cells)
